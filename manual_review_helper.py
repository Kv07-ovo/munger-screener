import csv
import shutil
import sys
from datetime import datetime
from pathlib import Path

# v2.1.0：Windows 控制台默认非 UTF-8（如 GBK），打印中文会报错，先重配。
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE_DIR = Path(__file__).parent
INPUT_PATH = BASE_DIR / "data" / "stocks.csv"
OUTPUT_PATH = BASE_DIR / "output" / "manual_review_needed.csv"
TEMPLATE_PATH = BASE_DIR / "output" / "manual_fill_template.csv"
XLSX_PATH = BASE_DIR / "output" / "manual_fill_template.xlsx"

# ── check-mode: which fields to scan ──────────────────────────────────────────

TEXT_FIELDS = [
    "industry",
    "circle_of_competence",
    "moat_reason",
    "management_reason",
    "risk_note",
    "risk_reason",
]

NUMERIC_FIELDS = [
    "brand_score",
    "switching_cost_score",
    "network_effect_score",
    "scale_advantage_score",
    "pricing_power_score",
    "moat_durability_score",
    "management_score",
    "confidence_score",
    "pe",
    "fcf_yield",
]

ALL_CHECKED = TEXT_FIELDS + NUMERIC_FIELDS
TEXT_SENTINEL = {"", "unknown", "manual_pending"}

# ── Industry exemptions for check_row ────────────────────────────────────────
_TICKER_INDUSTRY_HINTS: dict[str, str] = {
    "JPM": "银行", "BAC": "银行", "WFC": "银行", "GS": "银行", "C": "银行",
    "BRK-B": "综合控股", "BRK-A": "综合控股",
    "AIG": "保险", "MET": "保险", "PRU": "保险",
}

def _effective_industry(row: dict) -> str:
    ind = str(row.get("industry", "")).strip()
    if not ind or ind.lower() == "unknown":
        return _TICKER_INDUSTRY_HINTS.get(
            str(row.get("ticker", "")).strip().upper(), "")
    return ind

def _exempt_missing_fields(row: dict) -> set:
    """Fields to skip in missing-check for special industries."""
    ind = _effective_industry(row)
    exempt: set = set()
    if "银行" in ind:
        exempt |= {"fcf_yield", "gross_margin_5y_avg", "fcf_positive_years"}
    if "保险" in ind or "综合控股" in ind:
        exempt |= {"gross_margin_5y_avg", "fcf_yield"}
    return exempt

# ── template column order (English, canonical) ────────────────────────────────

TEMPLATE_FIELDS = [
    "ticker",
    "name",
    "industry",
    "brand_score",
    "switching_cost_score",
    "network_effect_score",
    "scale_advantage_score",
    "pricing_power_score",
    "moat_durability_score",
    "management_score",
    "circle_of_competence",
    "confidence_score",
    "risk_note",
    "moat_reason",
    "management_reason",
    "risk_reason",
]

# ── Chinese display labels (xlsx only, never written to stocks.csv) ────────────

FIELD_LABELS = {
    "ticker":                "股票代码",
    "name":                  "公司名称",
    "industry":              "行业",
    "brand_score":           "品牌强度",
    "switching_cost_score":  "转换成本",
    "network_effect_score":  "网络效应",
    "scale_advantage_score": "规模优势",
    "pricing_power_score":   "定价权",
    "moat_durability_score": "护城河持续性",
    "management_score":      "管理层评分",
    "circle_of_competence":  "能力圈",
    "confidence_score":      "数据可信度",
    "risk_note":             "风险短标签",
    "moat_reason":           "护城河理由",
    "management_reason":     "管理层理由",
    "risk_reason":           "风险详细说明",
}

# ── xlsx style constants ───────────────────────────────────────────────────────

FIELD_GROUPS = {
    "ticker":                "基础信息",
    "name":                  "基础信息",
    "industry":              "基础信息",
    "brand_score":           "护城河评分",
    "switching_cost_score":  "护城河评分",
    "network_effect_score":  "护城河评分",
    "scale_advantage_score": "护城河评分",
    "pricing_power_score":   "护城河评分",
    "moat_durability_score": "护城河评分",
    "management_score":      "管理与能力圈",
    "circle_of_competence":  "管理与能力圈",
    "confidence_score":      "管理与能力圈",
    "risk_note":             "风险与理由",
    "moat_reason":           "风险与理由",
    "management_reason":     "风险与理由",
    "risk_reason":           "风险与理由",
}

# ARGB colors (fully opaque = "FF" prefix)
HEADER_FILL = {
    "基础信息":    "FF9DC3E6",
    "护城河评分":  "FFA9D18E",
    "管理与能力圈": "FFFFD966",
    "风险与理由":  "FFF4B183",
}
DATA_FILL = {
    "基础信息":    "FFBDD7EE",
    "护城河评分":  "FFE2EFDA",
    "管理与能力圈": "FFFFF2CC",
    "风险与理由":  "FFFCE4D6",
}

WRAP_FIELDS = {"moat_reason", "management_reason", "risk_reason"}

SCORE_FIELDS = {
    "brand_score", "switching_cost_score", "network_effect_score",
    "scale_advantage_score", "pricing_power_score", "moat_durability_score",
    "management_score", "confidence_score",
}

RISK_NOTE_OPTIONS = [
    "high valuation", "growth slowdown", "regulatory risk", "competition risk",
    "cyclical risk", "margin pressure", "china exposure", "litigation risk",
    "execution risk", "debt risk", "data quality risk",
]

# (Chinese label, description shown in 字段说明 sheet)
FIELD_DESCRIPTIONS = [
    ("股票代码",    "股票 ticker，比如 AAPL、MSFT、GOOGL"),
    ("公司名称",    "公司中文名或英文名"),
    ("行业",        "公司所属行业，比如 科技、金融科技、消费品、银行"),
    ("品牌强度",    "品牌是否强，0-10 分"),
    ("转换成本",    "用户换到竞争对手的难度，0-10 分"),
    ("网络效应",    "用户越多，产品越有价值的程度，0-10 分"),
    ("规模优势",    "规模带来的成本、渠道、数据或供应链优势，0-10 分"),
    ("定价权",      "公司涨价后客户是否仍愿意买，0-10 分"),
    ("护城河持续性","竞争优势能维持多久，0-10 分"),
    ("管理层评分",  "管理层是否靠谱、是否善于资本配置，0-10 分"),
    ("能力圈",
     "inside  = 能力圈内，我比较懂\n"
     "edge    = 能力圈边缘，需要谨慎\n"
     "outside = 超出能力圈，暂时不建议研究"),
    ("数据可信度",  "你对数据和判断的信心，0-10 分"),
    ("风险短标签",
     "简短风险标签（可自由输入）\n"
     "high valuation   = 估值偏高\n"
     "growth slowdown  = 增长放缓\n"
     "regulatory risk  = 监管风险\n"
     "competition risk = 竞争风险\n"
     "cyclical risk    = 周期风险\n"
     "margin pressure  = 利润率压力\n"
     "china exposure   = 中国市场风险\n"
     "litigation risk  = 诉讼风险\n"
     "execution risk   = 执行风险\n"
     "debt risk        = 债务风险\n"
     "data quality risk= 数据质量风险"),
    ("护城河理由",  "为什么你认为它有护城河"),
    ("管理层理由",  "为什么你认为管理层好或不好"),
    ("风险详细说明","这家公司最主要的风险是什么"),
]

# xlsx row layout constants
_ROW_TITLE      = 1
_ROW_HINT       = 2
_ROW_CN_HEADERS = 3
_ROW_EN_FIELDS  = 4   # hidden; used by --apply-xlsx-template
_DATA_START     = 5

# ── helpers ────────────────────────────────────────────────────────────────────

def _is_missing_text(value: str) -> bool:
    return value.strip().lower() in TEXT_SENTINEL


def _is_missing_numeric(value: str) -> bool:
    v = value.strip()
    if not v:
        return True
    try:
        return float(v) == 0.0
    except ValueError:
        return v.lower() in TEXT_SENTINEL


def check_row(row: dict) -> list[str]:
    exempt = _exempt_missing_fields(row)
    missing = []
    for field in TEXT_FIELDS:
        if field not in exempt and _is_missing_text(row.get(field, "")):
            missing.append(field)
    for field in NUMERIC_FIELDS:
        if field not in exempt and _is_missing_numeric(row.get(field, "")):
            missing.append(field)
    return missing


def _estimate_col_width(field: str, rows: list[dict]) -> float:
    """Return a sensible column width capped at 40."""
    cn_label = FIELD_LABELS.get(field, field)
    # Chinese chars count as ~2 display units
    cn_display = sum(2 if ord(c) > 127 else 1 for c in cn_label)
    max_len = max(len(field), cn_display)
    for row in rows:
        val = str(row.get(field, ""))
        if field in WRAP_FIELDS:
            max_len = max(max_len, min(len(val), 28))
        else:
            display_len = sum(2 if ord(c) > 127 else 1 for c in val)
            max_len = max(max_len, display_len)
    return min(max_len + 3, 40)


def _row_height(row: dict) -> float:
    """Taller rows for data with long wrap-text content."""
    for field in WRAP_FIELDS:
        val = row.get(field, "")
        if val and len(val) > 10:
            lines = max(1, len(val) // 22)
            return min(15 * lines + 10, 120)
    return 20


# ── main scan ─────────────────────────────────────────────────────────────────

def main():
    with open(INPUT_PATH, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    review_rows = []
    for row in rows:
        missing = check_row(row)
        if missing:
            review_rows.append(
                {
                    "ticker": row["ticker"],
                    "name": row["name"],
                    "missing_count": len(missing),
                    "missing_fields": "|".join(missing),
                    **{f: row.get(f, "") for f in ALL_CHECKED},
                }
            )

    output_cols = ["ticker", "name", "missing_count", "missing_fields"] + ALL_CHECKED
    with open(OUTPUT_PATH, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        writer.writerows(review_rows)

    total = len(rows)
    flagged = len(review_rows)
    print(f"Checked {total} stocks -> {flagged} need manual review -> {OUTPUT_PATH}")
    print()
    if review_rows:
        print(f"{'Ticker':<10} {'Name':<12} {'Missing (#)':<12} Fields")
        print("-" * 80)
        for r in review_rows:
            print(f"{r['ticker']:<10} {r['name']:<12} {r['missing_count']:<12} {r['missing_fields']}")
    else:
        print("All stocks have complete manual fields.")


# ── xlsx generator ─────────────────────────────────────────────────────────────

def _generate_xlsx(template_rows: list[dict]) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("WARNING: openpyxl not installed. Skipping xlsx generation.")
        print("         Run: pip install openpyxl")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "补录模板"

    num_cols     = len(TEMPLATE_FIELDS)
    last_col     = get_column_letter(num_cols)
    last_data_row = _DATA_START - 1 + len(template_rows)

    # ── Row 1: merged title ────────────────────────────────────────────────────
    ws.merge_cells(f"A{_ROW_TITLE}:{last_col}{_ROW_TITLE}")
    title = ws.cell(_ROW_TITLE, 1)
    title.value = "Munger Screener 人工补录模板"
    title.font = Font(bold=True, size=14, color="FF1F3864")
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.fill = PatternFill(start_color="FFD6E4F0", end_color="FFD6E4F0", fill_type="solid")
    ws.row_dimensions[_ROW_TITLE].height = 30

    # ── Row 2: hint text ───────────────────────────────────────────────────────
    ws.merge_cells(f"A{_ROW_HINT}:{last_col}{_ROW_HINT}")
    hint = ws.cell(_ROW_HINT, 1)
    hint.value = "请填写中文表头对应的内容，程序会根据隐藏英文字段名自动导入。"
    hint.font = Font(italic=True, size=10, color="FF595959")
    hint.alignment = Alignment(horizontal="center", vertical="center")
    hint.fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
    ws.row_dimensions[_ROW_HINT].height = 20

    # ── Row 3: Chinese headers (visible, colored) ──────────────────────────────
    for ci, field in enumerate(TEMPLATE_FIELDS, start=1):
        cell = ws.cell(_ROW_CN_HEADERS, ci)
        cell.value = FIELD_LABELS.get(field, field)
        grp = FIELD_GROUPS.get(field, "")
        hc = HEADER_FILL.get(grp, "FFD9D9D9")
        cell.fill = PatternFill(start_color=hc, end_color=hc, fill_type="solid")
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[_ROW_CN_HEADERS].height = 36

    # ── Row 4: English field names (hidden; read by --apply-xlsx-template) ─────
    for ci, field in enumerate(TEMPLATE_FIELDS, start=1):
        cell = ws.cell(_ROW_EN_FIELDS, ci)
        cell.value = field
        cell.font = Font(size=9, color="FF808080")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[_ROW_EN_FIELDS].hidden = True

    # ── Data rows (row 5+) ─────────────────────────────────────────────────────
    for ri, row_data in enumerate(template_rows, start=_DATA_START):
        ws.row_dimensions[ri].height = _row_height(row_data)
        for ci, field in enumerate(TEMPLATE_FIELDS, start=1):
            cell = ws.cell(ri, ci)
            val = row_data.get(field, "")
            if field in SCORE_FIELDS and val != "":
                try:
                    cell.value = int(float(val))
                except (ValueError, TypeError):
                    cell.value = val
            else:
                cell.value = val
            grp = FIELD_GROUPS.get(field, "")
            dc = DATA_FILL.get(grp, "FFFFFFFF")
            cell.fill = PatternFill(start_color=dc, end_color=dc, fill_type="solid")
            cell.alignment = Alignment(wrap_text=(field in WRAP_FIELDS), vertical="top")

    # ── Freeze: rows 1-4 (3 visible + 1 hidden) + cols A-C  →  D5 ─────────────
    ws.freeze_panes = f"D{_DATA_START}"

    # ── Auto filter on Chinese header row ──────────────────────────────────────
    ws.auto_filter.ref = f"A{_ROW_CN_HEADERS}:{last_col}{last_data_row}"

    # ── Column widths ──────────────────────────────────────────────────────────
    for ci, field in enumerate(TEMPLATE_FIELDS, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = _estimate_col_width(field, template_rows)

    # ── Data validation: score fields 0-10 ────────────────────────────────────
    score_formula = '"0,1,2,3,4,5,6,7,8,9,10"'
    for ci, field in enumerate(TEMPLATE_FIELDS, start=1):
        if field in SCORE_FIELDS:
            col = get_column_letter(ci)
            dv = DataValidation(type="list", formula1=score_formula, allow_blank=True)
            dv.sqref = f"{col}{_DATA_START}:{col}{last_data_row}"
            ws.add_data_validation(dv)

    # ── Data validation: circle_of_competence ─────────────────────────────────
    coc_ci  = TEMPLATE_FIELDS.index("circle_of_competence") + 1
    coc_col = get_column_letter(coc_ci)
    dv_coc  = DataValidation(type="list", formula1='"inside,edge,outside"', allow_blank=True)
    dv_coc.sqref = f"{coc_col}{_DATA_START}:{coc_col}{last_data_row}"
    ws.add_data_validation(dv_coc)

    # ── Data validation: risk_note (dropdown + free input allowed) ─────────────
    rn_ci  = TEMPLATE_FIELDS.index("risk_note") + 1
    rn_col = get_column_letter(rn_ci)
    dv_rn  = DataValidation(
        type="list",
        formula1='"' + ",".join(RISK_NOTE_OPTIONS) + '"',
        allow_blank=True,
    )
    dv_rn.showErrorMessage = False   # allows typing any text beyond the list
    dv_rn.sqref = f"{rn_col}{_DATA_START}:{rn_col}{last_data_row}"
    ws.add_data_validation(dv_rn)

    # ── Sheet 2: 字段说明 ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("字段说明")
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 52

    for ci, hdr in enumerate(["字段", "说明"], start=1):
        cell = ws2.cell(1, ci)
        cell.value = hdr
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(start_color="FFD6E4F0", end_color="FFD6E4F0", fill_type="solid")
    ws2.row_dimensions[1].height = 24

    for ri, (cn_label, desc) in enumerate(FIELD_DESCRIPTIONS, start=2):
        name_cell = ws2.cell(ri, 1)
        name_cell.value = cn_label
        name_cell.font = Font(bold=True, size=10)
        name_cell.alignment = Alignment(vertical="top")
        desc_cell = ws2.cell(ri, 2)
        desc_cell.value = desc
        desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if "\n" in desc:
            ws2.row_dimensions[ri].height = 15 * (desc.count("\n") + 1) + 5

    ws2.freeze_panes = "A2"

    wb.save(XLSX_PATH)


# ── template generation ────────────────────────────────────────────────────────

def generate_template():
    if not OUTPUT_PATH.exists():
        print(f"ERROR: {OUTPUT_PATH} not found. Run without --template first.")
        sys.exit(1)

    with open(OUTPUT_PATH, encoding="utf-8-sig") as f:
        review_rows = list(csv.DictReader(f))

    if not review_rows:
        print("No stocks need manual review. Template not generated.")
        return

    with open(INPUT_PATH, encoding="utf-8-sig") as f:
        stocks = {row["ticker"]: row for row in csv.DictReader(f)}

    template_rows = []
    for r in review_rows:
        ticker = r["ticker"]
        src = stocks.get(ticker, {})
        missing_set = set(r["missing_fields"].split("|")) if r["missing_fields"] else set()
        out = {}
        for field in TEMPLATE_FIELDS:
            if field in ("ticker", "name"):
                out[field] = r[field]
            elif field in missing_set:
                out[field] = ""
            else:
                out[field] = src.get(field, "")
        if any(out[f] == "" for f in TEMPLATE_FIELDS if f not in ("ticker", "name")):
            template_rows.append(out)

    # CSV keeps English field names (program-compatible)
    with open(TEMPLATE_PATH, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=TEMPLATE_FIELDS)
        writer.writeheader()
        writer.writerows(template_rows)

    # XLSX gets Chinese display headers
    _generate_xlsx(template_rows)

    print("已生成：")
    print(f"  * {TEMPLATE_PATH}")
    print(f"  * {XLSX_PATH}")
    print()
    print(f"  共 {len(template_rows)} 只股票需要补录：")
    for r in template_rows:
        blanks = [f for f in TEMPLATE_FIELDS if r[f] == "" and f not in ("ticker", "name")]
        print(f"  {r['ticker']:<10} {r['name']:<12} 待填字段: {', '.join(blanks)}")
    print()
    print("请优先打开 xlsx 文件填写。")
    print("填写完成后，运行：python manual_review_helper.py --apply-xlsx-template")


# ── xlsx import ───────────────────────────────────────────────────────────────

def apply_xlsx_template():
    """Read the filled xlsx template and write non-empty values back to stocks.csv."""
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl not installed.  Run: pip install openpyxl")
        sys.exit(1)

    if not XLSX_PATH.exists():
        print(f"ERROR: {XLSX_PATH} not found.")
        print("       Generate it first: python manual_review_helper.py --template")
        sys.exit(1)

    if not INPUT_PATH.exists():
        print(f"ERROR: {INPUT_PATH} not found.")
        sys.exit(1)

    # ── 1. Read xlsx ──────────────────────────────────────────────────────────
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    sheet_name = "补录模板"
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: Sheet '{sheet_name}' not found in {XLSX_PATH}")
        sys.exit(1)
    ws = wb[sheet_name]

    def _read_row_headers(row_num: int) -> list[str]:
        result = []
        for ci in range(1, ws.max_column + 1):
            raw = ws.cell(row_num, ci).value
            if raw is None:
                break
            result.append(str(raw).strip())
        return result

    # Format detection:
    #   New format: English field names in row 4 (hidden), data from row 5
    #   Old format: English field names in row 2, data from row 3
    en_headers_r4 = _read_row_headers(_ROW_EN_FIELDS)
    if "ticker" in en_headers_r4:
        xlsx_headers  = en_headers_r4
        data_start    = _DATA_START
    else:
        en_headers_r2 = _read_row_headers(2)
        if "ticker" in en_headers_r2:
            xlsx_headers  = en_headers_r2
            data_start    = 3
        else:
            print("ERROR: Cannot find English field names ('ticker') in xlsx.")
            print("       Expected in row 4 (new format) or row 2 (old format).")
            sys.exit(1)

    ticker_ci = xlsx_headers.index("ticker") + 1

    def _cell_str(raw) -> str:
        if raw is None:
            return ""
        if isinstance(raw, float) and raw == int(raw):
            return str(int(raw))
        return str(raw).strip()

    xlsx_data: dict[str, dict] = {}
    for ri in range(data_start, ws.max_row + 1):
        ticker_raw = ws.cell(ri, ticker_ci).value
        if ticker_raw is None:
            continue
        ticker = _cell_str(ticker_raw)
        if not ticker:
            continue
        xlsx_data[ticker] = {
            hdr: _cell_str(ws.cell(ri, ci + 1).value)
            for ci, hdr in enumerate(xlsx_headers)
        }

    if not xlsx_data:
        print("No data rows found in xlsx template. Nothing to apply.")
        return

    # ── 2. Read stocks.csv ────────────────────────────────────────────────────
    with open(INPUT_PATH, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        orig_fieldnames = list(reader.fieldnames)
        stocks: list[dict] = list(reader)

    stocks_index = {row["ticker"]: row for row in stocks}

    # ── 3. Backup ─────────────────────────────────────────────────────────────
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = BASE_DIR / "data" / f"stocks_backup_{ts}.csv"
    shutil.copy2(INPUT_PATH, backup_path)
    print(f"备份已创建: {backup_path}")
    print()

    # ── 4. Apply updates (skip ticker; update all other non-empty fields) ─────
    updatable = [f for f in xlsx_headers if f != "ticker"]

    not_found: list[str] = []
    summary: dict[str, dict] = {}

    for ticker, xlsx_row in xlsx_data.items():
        if ticker not in stocks_index:
            not_found.append(ticker)
            continue

        stock_row = stocks_index[ticker]
        updated, skipped = [], []

        for field in updatable:
            val = xlsx_row.get(field, "")
            if val == "":
                skipped.append(field)
            elif field in orig_fieldnames:
                stock_row[field] = val
                updated.append(field)

        summary[ticker] = {
            "name":    xlsx_row.get("name", stock_row.get("name", "")),
            "updated": updated,
            "skipped": skipped,
        }

    # ── 5. Write back stocks.csv ──────────────────────────────────────────────
    with open(INPUT_PATH, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=orig_fieldnames)
        writer.writeheader()
        writer.writerows(stocks)

    # ── 6. Print summary ──────────────────────────────────────────────────────
    print("=" * 70)
    print("导入结果")
    print("=" * 70)

    total_stocks = 0
    total_fields = 0
    for ticker, info in summary.items():
        updated = info["updated"]
        skipped = info["skipped"]
        name    = info["name"]
        if updated:
            total_stocks += 1
            total_fields += len(updated)
            print(f"\n  {ticker} ({name})")
            print(f"    更新字段 ({len(updated)}): {', '.join(updated)}")
            if skipped:
                print(f"    跳过空字段 ({len(skipped)}): {', '.join(skipped)}")
        else:
            print(f"\n  {ticker} ({name})  — 无更新（模板中所有字段均为空）")

    if not_found:
        print(f"\n  未找到的 ticker（不在 stocks.csv 中）: {', '.join(not_found)}")

    print()
    print(f"共更新 {total_stocks} 只股票，{total_fields} 个字段")
    print(f"stocks.csv 已写回: {INPUT_PATH}")

    # ── 7. Re-run check ───────────────────────────────────────────────────────
    print()
    print("─" * 70)
    print("重新检查人工字段完整性 ...")
    print("─" * 70)
    main()


# ── entry point ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    if "--template" in sys.argv:
        generate_template()
    elif "--apply-xlsx-template" in sys.argv:
        apply_xlsx_template()
    else:
        main()
